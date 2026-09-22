"""
Project 4 - Build the weekly Distribution Ops report pack (Excel).

Builds the formatted Excel, only if check passes - The automation

"""
import os
import re
import sys
from datetime import date
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
DATA = os.path.join(ROOT, "data")
AS_OF = "2026-06-28"


# connections
def local_connection():
    import duckdb
    con = duckdb.connect()
    for s in ("RAW", "MART", "RPT"):
        con.execute(f"CREATE SCHEMA {s}")
    for t in ["dim_date", "dim_product", "dim_distributor", "fact_orders", "fact_shipments",
              "fact_billing", "fact_pos", "fact_dist_inventory"]:
        con.execute(f"CREATE TABLE RAW.{t.upper()} AS SELECT * FROM read_csv_auto('{DATA}/{t}.csv')")
    # Snowflake -> DuckDB shims
    con.execute("CREATE MACRO IFF(c, a, b) AS CASE WHEN c THEN a ELSE b END")
    con.execute("""CREATE MACRO DATEADD(u, n, d) AS CASE lower(u)
                     WHEN 'day'     THEN CAST(CAST(d AS DATE) + CAST(n AS INT) AS DATE)
                     WHEN 'week'    THEN CAST(CAST(d AS DATE) + 7 * CAST(n AS INT) AS DATE)
                     WHEN 'month'   THEN CAST(CAST(d AS DATE) + to_months(CAST(n AS INT)) AS DATE)
                     WHEN 'quarter' THEN CAST(CAST(d AS DATE) + to_months(3 * CAST(n AS INT)) AS DATE)
                   END""")

    def run_file(path):
        sql = re.sub(r"(?im)^\s*USE [^;]*;", "", open(path).read())
        for stmt in sql.split(";"):
            if re.search(r"(?im)^\s*(CREATE|SELECT|WITH)", stmt) and "CREATE WAREHOUSE" not in stmt:
                if re.search(r"(?im)^\s*CREATE (SCHEMA|DATABASE)", stmt):
                    continue
                con.execute(stmt)

    run_file(os.path.join(ROOT, "project1_channel_health", "02_mart_views.sql"))
    run_file(os.path.join(HERE, "01_report_views.sql"))
    return lambda q: con.execute(q).df()


def snowflake_connection():
    import snowflake.connector
    con = snowflake.connector.connect(account=os.environ["SF_ACCOUNT"], user=os.environ["SF_USER"],
                                      password=os.environ["SF_PASSWORD"], warehouse="ANALYTICS_WH",
                                      database="CHANNEL_ANALYTICS")
    def q(sql):
        cur = con.cursor()
        cur.execute(sql)
        df = cur.fetch_pandas_all()
        cur.close()
        return df
    return q


# build
def build(query):
    checks = query("SELECT * FROM RPT.VW_REPORT_CHECKS")
    checks.columns = [c.upper() for c in checks.columns]
    print(checks.to_string(index=False))
    if (checks.RESULT != "PASS").any():
        sys.exit("\nTie-out check FAILED - report pack not published. Investigate before sending.")

    latest = query("SELECT MAX(WEEK_START) AS W FROM RPT.VW_BACKLOG_WATERFALL_WEEKLY").iloc[0, 0]
    latest = pd.Timestamp(latest).date()

    waterfall = query(f"""
        SELECT w.WEEK_START, w.PRODUCT_FAMILY,
               SUM(BEGIN_BACKLOG) AS BEGIN_BACKLOG, SUM(BOOKINGS) AS BOOKINGS,
               SUM(SHIPMENTS) AS SHIPMENTS, SUM(END_BACKLOG) AS END_BACKLOG, SUM(NET_CHANGE) AS NET_CHANGE
        FROM RPT.VW_BACKLOG_WATERFALL_WEEKLY w
        WHERE w.WEEK_START >= DATEADD('week', -12, '{latest}'::DATE)
        GROUP BY 1, 2 ORDER BY 1 DESC, 2""")
    delinq = query(f"""
        SELECT WEEK_START, PRODUCT_FAMILY, SUM(DELINQUENT_LINES) AS DELINQUENT_LINES,
               SUM(DELINQUENT_VALUE) AS DELINQUENT_VALUE
        FROM RPT.VW_DELINQUENCY_WEEKLY
        WHERE WEEK_START >= DATEADD('week', -12, '{latest}'::DATE)
        GROUP BY 1, 2 ORDER BY 1 DESC, 4 DESC""")
    pace = query("SELECT * FROM RPT.VW_QTD_BILLINGS_PACE WHERE QTR >= '2025-01-01' ORDER BY QTR DESC, WEEK_OF_QTR")
    otd = query("""
        SELECT COMMIT_MONTH, PRODUCT_FAMILY, SUM(LINES_DUE) AS LINES_DUE, SUM(LINES_ON_TIME) AS LINES_ON_TIME,
               SUM(LINES_ON_TIME) / SUM(LINES_DUE) AS OTD_TO_COMMIT
        FROM RPT.VW_OTD_MONTHLY WHERE COMMIT_MONTH >= '2025-07-01'
        GROUP BY 1, 2 ORDER BY 1 DESC, 2""")
    score = query("""
        SELECT MTH, DISTRIBUTOR_NAME, REGION, BILLINGS, SHARE_OF_BILLINGS, BILLINGS_MOM_PCT, BILLINGS_RANK,
               RANK_CHANGE, OTD_TO_COMMIT, WEEKS_OF_INVENTORY, DELINQUENT_VALUE, HEALTH_STATUS
        FROM RPT.VW_DISTRIBUTOR_SCORECARD
        WHERE MTH = (SELECT MAX(MTH) FROM RPT.VW_DISTRIBUTOR_SCORECARD)
        ORDER BY BILLINGS_RANK""")
    for df in (waterfall, delinq, pace, otd, score):
        df.columns = [c.upper() for c in df.columns]

    # headline KPIs
    wk = waterfall[pd.to_datetime(waterfall.WEEK_START).dt.date == latest]
    dl = delinq[pd.to_datetime(delinq.WEEK_START).dt.date == latest]
    cur_q = pace[pace.QTR == pace.QTR.max()].sort_values("WEEK_OF_QTR").iloc[-1]
    last_m = otd[otd.COMMIT_MONTH == otd.COMMIT_MONTH.max()]
    kpis = pd.DataFrame([
        ("Week ending", str(pd.Timestamp(latest) + pd.Timedelta(days=6))[:10], ""),
        ("Open backlog ($)", wk.END_BACKLOG.sum(), f"{wk.NET_CHANGE.sum():+,.0f} vs prior week"),
        ("Bookings this week ($)", wk.BOOKINGS.sum(), ""),
        ("Shipments this week ($)", wk.SHIPMENTS.sum(), ""),
        ("Book-to-ship ratio (week)", wk.BOOKINGS.sum() / wk.SHIPMENTS.sum(), ">1 = backlog growing"),
        ("Delinquent backlog ($)", dl.DELINQUENT_VALUE.sum(),
         f"{dl.sort_values('DELINQUENT_VALUE').iloc[-1].PRODUCT_FAMILY} = "
         f"{dl.DELINQUENT_VALUE.max() / dl.DELINQUENT_VALUE.sum():.0%} of total"),
        ("QTD billings ($)", cur_q.QTD_BILLINGS,
         f"{cur_q.PACE_VS_PRIOR_QTR:+.1%} vs prior quarter at week {int(cur_q.WEEK_OF_QTR)}"),
        ("OTD to commit (latest month)", last_m.LINES_ON_TIME.sum() / last_m.LINES_DUE.sum(), ""),
        ("Distributors flagged Red", int((score.HEALTH_STATUS == "Red").sum()),
         ", ".join(score[score.HEALTH_STATUS == "Red"].DISTRIBUTOR_NAME)),
    ], columns=["KPI", "Value", "Comment"])

    out = os.path.join(HERE, f"Distribution_Ops_Weekly_Pack_{latest}.xlsx")
    sheets = {"Summary": kpis, "Backlog Waterfall": waterfall, "Delinquency": delinq,
              "QTD Billings Pace": pace, "On-Time Delivery": otd, "Distributor Scorecard": score,
              "Tie-out Checks": checks}
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    format_workbook(out)
    print(f"\nPublished -> {out}")


# formatting
MONEY = {"BEGIN_BACKLOG", "BOOKINGS", "SHIPMENTS", "END_BACKLOG", "NET_CHANGE", "DELINQUENT_VALUE",
         "BILLINGS", "QTD_BILLINGS", "PRIOR_QTR_QTD_SAME_WEEK", "REPORT_VALUE", "SOURCE_VALUE"}
PCT = {"PACE_VS_PRIOR_QTR", "PCT_OF_QTR_BILLED", "OTD_TO_COMMIT", "SHARE_OF_BILLINGS", "BILLINGS_MOM_PCT"}


def format_workbook(path):
    wb = load_workbook(path)
    head_fill = PatternFill("solid", start_color="1F3864")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        headers = [c.value for c in ws[1]]
        for col_idx, h in enumerate(headers, start=1):
            letter = ws.cell(1, col_idx).column_letter
            ws.column_dimensions[letter].width = max(14, min(45, len(str(h)) + 4))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell = row[0]
                cell.font = Font(name="Arial")
                if h in MONEY:
                    cell.number_format = '$#,##0;($#,##0);-'
                elif h in PCT:
                    cell.number_format = '0.0%'
                elif h in ("WEEKS_OF_INVENTORY", "AVG_DAYS_PAST_DUE"):
                    cell.number_format = '0.0'
                elif isinstance(cell.value, (date, pd.Timestamp)):
                    cell.number_format = 'yyyy-mm-dd'
        ws.freeze_panes = "A2"
    s = wb["Summary"]
    for r, fmt in zip(range(3, 11), ['$#,##0', '$#,##0', '$#,##0', '0.00', '$#,##0', '$#,##0', '0.0%', '0']):
        s.cell(r, 2).number_format = fmt
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 32, 18, 55
    sc = wb["Distributor Scorecard"]
    col = [c.value for c in sc[1]].index("HEALTH_STATUS") + 1
    rng = f"{sc.cell(2, col).coordinate}:{sc.cell(sc.max_row, col).coordinate}"
    for val, color in (("Red", "F8CBAD"), ("Amber", "FFE699"), ("Green", "C6EFCE")):
        sc.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                      fill=PatternFill("solid", start_color=color)))
    wb.save(path)


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "--local"
    build(local_connection() if mode == "--local" else snowflake_connection())
